#!/usr/bin/env python3
"""Source-traceable study pipeline: inventory, extract, matrix, qa, export.

Local only (stdlib + optional pdfplumber / python-docx / pandoc / soffice).
Workspace layout under --root:
  sources/     original files (never modified)
  extracted/   <SID>.jsonl  one record per text block / table / equation
               tables/<SID>_<locator>_<n>.csv
  ledger/      inventory.json, evidence.csv, reports/*.json
  drafts/      *.md chapters
  export/      md / docx / pdf / html
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".md", ".txt", ".csv"}
EVIDENCE_FIELDS = ["claim_id", "claim", "source_id", "locator", "quote", "support", "confidence", "note"]
SUPPORT_LEVELS = {"direct", "derived", "inferred"}
CITE_RE = re.compile(r"\[@(S\d{3})(?:\s+([^\]]+))?\]")
CLAIM_RE = re.compile(r"\[(C\d{3,})\]")
EQ_LINE_RE = re.compile(r"^(?=.*=)(?=.*[A-Za-z0-9])[^=]{1,80}=[^=].{0,120}$")
SENT_SPLIT = re.compile(r"(?<=[.!?。！？])\s+")
NUM_RE = re.compile(r"\d")


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace(" ", " ")).strip().lower()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def paths(root: Path) -> dict[str, Path]:
    p = {k: root / k for k in ("sources", "extracted", "ledger", "drafts", "export")}
    p["inventory"] = p["ledger"] / "inventory.json"
    p["evidence"] = p["ledger"] / "evidence.csv"
    p["reports"] = p["ledger"] / "reports"
    return p


def load_inventory(root: Path) -> dict:
    inv = paths(root)["inventory"]
    if not inv.exists():
        sys.exit(f"missing {inv}; run `inventory` first")
    return json.loads(inv.read_text(encoding="utf-8"))


def load_blocks(root: Path, sid: str) -> list[dict]:
    f = paths(root)["extracted"] / f"{sid}.jsonl"
    if not f.exists():
        return []
    return [json.loads(line) for line in f.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_report(root: Path, name: str, report: dict) -> None:
    rd = paths(root)["reports"]
    rd.mkdir(parents=True, exist_ok=True)
    (rd / f"{name}.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- inventory
def cmd_init(a):
    for k, p in paths(Path(a.root)).items():
        if k in ("sources", "extracted", "ledger", "drafts", "export", "reports"):
            p.mkdir(parents=True, exist_ok=True)
    print(f"initialized {a.root}")


def cmd_inventory(a):
    root = Path(a.root)
    P = paths(root)
    P["ledger"].mkdir(parents=True, exist_ok=True)
    old = json.loads(P["inventory"].read_text()) if P["inventory"].exists() else {"sources": {}}
    by_hash = {s["sha256"]: sid for sid, s in old["sources"].items()}
    sources = dict(old["sources"])
    next_n = max([int(k[1:]) for k in sources] + [0]) + 1
    skipped = []
    for f in sorted(P["sources"].rglob("*")):
        if not f.is_file():
            continue
        if f.suffix.lower() not in SUPPORTED:
            skipped.append(str(f.relative_to(root)))
            continue
        h = sha256(f)
        rel = str(f.relative_to(root))
        if h in by_hash:
            sources[by_hash[h]].setdefault("duplicate_paths", [])
            if rel != sources[by_hash[h]]["path"] and rel not in sources[by_hash[h]]["duplicate_paths"]:
                sources[by_hash[h]]["duplicate_paths"].append(rel)
            continue
        sid = f"S{next_n:03d}"
        next_n += 1
        sources[sid] = {"path": rel, "title": f.stem.replace("_", " "), "type": f.suffix.lower().lstrip("."),
                        "bytes": f.stat().st_size, "sha256": h, "extracted": False,
                        "reliability": "unrated", "notes": ""}
        by_hash[h] = sid
    P["inventory"].write_text(json.dumps({"sources": sources}, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"{len(sources)} sources; skipped unsupported: {skipped or 'none'}")


# ------------------------------------------------------------------ extract
def _equations(text: str) -> list[str]:
    out = []
    for line in text.splitlines():
        line = line.strip()
        if EQ_LINE_RE.match(line) and len(line.split()) <= 20 and not line.endswith(":"):
            out.append(line)
    return out


def _extract_pdf(path: Path, sid: str, tdir: Path) -> list[dict]:
    try:
        import pdfplumber
    except ImportError:
        sys.exit("pdfplumber missing: use the liteparse or markitdown skill for this PDF")
    blocks = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            loc = f"p.{i}"
            text = page.extract_text() or ""
            blocks.append({"source_id": sid, "locator": loc, "kind": "text", "text": text})
            for eq in _equations(text):
                blocks.append({"source_id": sid, "locator": loc, "kind": "equation", "text": eq})
            for n, tbl in enumerate(page.extract_tables(), 1):
                rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in tbl]
                fn = tdir / f"{sid}_p{i}_t{n}.csv"
                with fn.open("w", newline="", encoding="utf-8") as fh:
                    csv.writer(fh).writerows(rows)
                blocks.append({"source_id": sid, "locator": f"{loc} table {n}", "kind": "table",
                               "text": "\n".join(" | ".join(r) for r in rows), "csv": str(fn.name)})
            if not text.strip():
                blocks.append({"source_id": sid, "locator": loc, "kind": "warning",
                               "text": "no text layer; OCR with the liteparse skill"})
    return blocks


def _extract_docx(path: Path, sid: str, tdir: Path) -> list[dict]:
    import docx
    d = docx.Document(str(path))
    blocks = []
    for i, p in enumerate(d.paragraphs, 1):
        if p.text.strip():
            blocks.append({"source_id": sid, "locator": f"para {i}", "kind": "text", "text": p.text})
            for eq in _equations(p.text):
                blocks.append({"source_id": sid, "locator": f"para {i}", "kind": "equation", "text": eq})
    for n, t in enumerate(d.tables, 1):
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        fn = tdir / f"{sid}_t{n}.csv"
        with fn.open("w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows(rows)
        blocks.append({"source_id": sid, "locator": f"table {n}", "kind": "table",
                       "text": "\n".join(" | ".join(r) for r in rows), "csv": fn.name})
    return blocks


def _extract_text(path: Path, sid: str) -> list[dict]:
    blocks = []
    for i, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
        if line.strip():
            blocks.append({"source_id": sid, "locator": f"line {i}", "kind": "text", "text": line})
            for eq in _equations(line):
                blocks.append({"source_id": sid, "locator": f"line {i}", "kind": "equation", "text": eq})
    return blocks


def cmd_extract(a):
    root = Path(a.root)
    P = paths(root)
    inv = load_inventory(root)
    tdir = P["extracted"] / "tables"
    tdir.mkdir(parents=True, exist_ok=True)
    for sid, s in inv["sources"].items():
        if s["extracted"] and not a.force:
            continue
        f = root / s["path"]
        if sha256(f) != s["sha256"]:
            print(f"WARN {sid}: file changed since inventory; re-run inventory", file=sys.stderr)
            continue
        ext = f.suffix.lower()
        if ext == ".pdf":
            blocks = _extract_pdf(f, sid, tdir)
        elif ext == ".docx":
            blocks = _extract_docx(f, sid, tdir)
        else:
            blocks = _extract_text(f, sid)
        with (P["extracted"] / f"{sid}.jsonl").open("w", encoding="utf-8") as fh:
            for b in blocks:
                fh.write(json.dumps(b, ensure_ascii=False) + "\n")
        s["extracted"] = True
        kinds = {k: sum(b["kind"] == k for b in blocks) for k in ("text", "table", "equation", "warning")}
        print(f"{sid} {s['path']}: {kinds}")
    P["inventory"].write_text(json.dumps(inv, indent=2, ensure_ascii=False), encoding="utf-8")


# ------------------------------------------------------------------- matrix
def _locator_text(blocks: list[dict], locator: str) -> str | None:
    want = norm(locator)
    txt = [b["text"] for b in blocks if norm(b["locator"]) == want or norm(b["locator"]).startswith(want + " ")]
    return "\n".join(txt) if txt else None


def cmd_matrix_init(a):
    ev = paths(Path(a.root))["evidence"]
    ev.parent.mkdir(parents=True, exist_ok=True)
    if ev.exists() and not a.force:
        sys.exit(f"{ev} exists; use --force")
    with ev.open("w", newline="", encoding="utf-8") as fh:
        csv.DictWriter(fh, EVIDENCE_FIELDS).writeheader()
    print(f"wrote {ev}")


def validate_matrix(root: Path) -> tuple[dict, dict]:
    """Return (report, ok_claims). A claim passes only if its source, locator and quote resolve."""
    inv = load_inventory(root)["sources"]
    ev = paths(root)["evidence"]
    errors, ok = [], {}
    seen = set()
    if not ev.exists():
        return {"errors": [f"missing {ev}"], "claims": 0}, {}
    with ev.open(newline="", encoding="utf-8") as fh:
        rdr = csv.DictReader(fh)
        missing = [c for c in EVIDENCE_FIELDS if c not in (rdr.fieldnames or [])]
        if missing:
            return {"errors": [f"missing columns {missing}"], "claims": 0}, {}
        rows = list(rdr)
    blocks_cache: dict[str, list[dict]] = {}
    for n, r in enumerate(rows, 2):
        cid = r["claim_id"].strip()
        errs = []
        if not re.fullmatch(r"C\d{3,}", cid):
            errs.append("claim_id must look like C001")
        if cid in seen:
            errs.append("duplicate claim_id")
        seen.add(cid)
        sid = r["source_id"].strip()
        if sid not in inv:
            errs.append(f"unknown source_id {sid!r}")
        if r["support"].strip() not in SUPPORT_LEVELS:
            errs.append(f"support must be one of {sorted(SUPPORT_LEVELS)}")
        if not r["claim"].strip():
            errs.append("empty claim")
        if sid in inv:
            blocks = blocks_cache.setdefault(sid, load_blocks(root, sid))
            if not blocks:
                errs.append("source not extracted")
            else:
                loc_text = _locator_text(blocks, r["locator"].strip())
                if loc_text is None:
                    errs.append(f"locator {r['locator']!r} not found in {sid}")
                elif r["support"].strip() == "direct":
                    q = norm(r["quote"])
                    if not q:
                        errs.append("direct support requires a verbatim quote")
                    elif q not in norm(loc_text):
                        errs.append("quote not found verbatim at locator")
                elif not r["note"].strip():
                    errs.append("derived/inferred claims need a note explaining the step")
        if errs:
            errors.append({"row": n, "claim_id": cid, "errors": errs})
        else:
            ok[cid] = r
    return {"claims": len(rows), "passed": len(ok), "errors": errors}, ok


def cmd_matrix(a):
    root = Path(a.root)
    rep, _ = validate_matrix(root)
    write_report(root, "matrix", rep)
    print(json.dumps(rep, indent=2, ensure_ascii=False))
    sys.exit(1 if rep.get("errors") else 0)


# ----------------------------------------------------------------------- qa
def cmd_qa(a):
    root = Path(a.root)
    inv = load_inventory(root)["sources"]
    _, ok = validate_matrix(root)
    cache: dict[str, list[dict]] = {}
    errors, warnings, used_sources = [], [], set()
    drafts = [Path(d) for d in a.drafts] or sorted(paths(root)["drafts"].glob("*.md"))
    for d in drafts:
        in_code = False
        for ln, line in enumerate(d.read_text(encoding="utf-8").splitlines(), 1):
            if line.strip().startswith("```"):
                in_code = not in_code
            if in_code or line.startswith("#") or not line.strip():
                continue
            for sid, loc in CITE_RE.findall(line):
                used_sources.add(sid)
                if sid not in inv:
                    errors.append(f"{d.name}:{ln} unknown source {sid}")
                    continue
                if not loc.strip():
                    errors.append(f"{d.name}:{ln} [@{sid}] has no locator")
                    continue
                blocks = cache.setdefault(sid, load_blocks(root, sid))
                if _locator_text(blocks, loc.strip()) is None:
                    errors.append(f"{d.name}:{ln} locator {loc!r} not found in {sid}")
            for cid in CLAIM_RE.findall(line):
                if cid not in ok:
                    errors.append(f"{d.name}:{ln} claim {cid} missing or failed matrix validation")
                else:
                    used_sources.add(ok[cid]["source_id"])
            cited = bool(CITE_RE.search(line) or CLAIM_RE.search(line))
            if not cited and "{own}" not in line:
                for sent in SENT_SPLIT.split(line):
                    if NUM_RE.search(sent) and len(sent.split()) > 4 and not sent.lstrip().startswith(("|", "-", "$")):
                        warnings.append(f"{d.name}:{ln} numeric statement without citation (mark {{own}} if self-derived): {sent[:70]}")
                if line.strip().startswith("$$") and "{own}" not in line:
                    warnings.append(f"{d.name}:{ln} equation without citation or {{own}} marker")
    unused = sorted(set(inv) - used_sources)
    rep = {"drafts": [d.name for d in drafts], "errors": errors, "warnings": warnings, "unused_sources": unused}
    write_report(root, "qa", rep)
    print(json.dumps(rep, indent=2, ensure_ascii=False))
    sys.exit(1 if errors or (a.strict and warnings) else 0)


# ------------------------------------------------------------------- export
def _references(root: Path, md_files: list[Path]) -> str:
    inv = load_inventory(root)["sources"]
    _, ok = validate_matrix(root)
    used = set()
    for f in md_files:
        text = f.read_text(encoding="utf-8")
        used |= {s for s, _ in CITE_RE.findall(text)}
        used |= {ok[c]["source_id"] for c in CLAIM_RE.findall(text) if c in ok}
    lines = ["# References", ""]
    for sid in sorted(used & set(inv)):
        s = inv[sid]
        lines.append(f"- **{sid}** {s['title']} (`{s['path']}`, sha256 {s['sha256'][:12]})")
    return "\n".join(lines) + "\n"


def cmd_export(a):
    root = Path(a.root)
    P = paths(root)
    P["export"].mkdir(parents=True, exist_ok=True)
    drafts = [Path(d) for d in a.drafts] or sorted(P["drafts"].glob("*.md"))
    if not drafts:
        sys.exit("no drafts")
    if not a.skip_qa:
        qa = argparse.Namespace(root=a.root, drafts=[str(d) for d in drafts], strict=False)
        try:
            cmd_qa(qa)
        except SystemExit as e:
            if e.code:
                sys.exit("export blocked: qa errors (see ledger/reports/qa.json). Use --skip-qa only for drafts.")
    name = a.name
    combined = P["export"] / f"{name}.md"
    combined.write_text("\n\n".join(d.read_text(encoding="utf-8") for d in drafts) + "\n\n" + _references(root, drafts),
                        encoding="utf-8")
    outs = [combined]
    pandoc = shutil.which("pandoc")
    for fmt in a.formats:
        if fmt == "md":
            continue
        if not pandoc:
            sys.exit("pandoc not found")
        if fmt in ("docx", "html"):
            out = P["export"] / f"{name}.{fmt}"
            cmd = [pandoc, str(combined), "-o", str(out), "--toc"] + (["--standalone"] if fmt == "html" else [])
            subprocess.run(cmd, check=True)
            outs.append(out)
        elif fmt == "pdf":
            docx = P["export"] / f"{name}.docx"
            if not docx.exists():
                subprocess.run([pandoc, str(combined), "-o", str(docx), "--toc"], check=True)
            office = shutil.which("soffice") or shutil.which("libreoffice")
            if not office:
                sys.exit("no soffice/libreoffice for pdf export")
            subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(P["export"]), str(docx)],
                           check=True, capture_output=True)
            outs.append(P["export"] / f"{name}.pdf")
    for o in outs:
        print(o, o.stat().st_size if o.exists() else "MISSING")
        if not o.exists():
            sys.exit(1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", default="study")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init").set_defaults(fn=cmd_init)
    sub.add_parser("inventory").set_defaults(fn=cmd_inventory)
    p = sub.add_parser("extract")
    p.add_argument("--force", action="store_true")
    p.set_defaults(fn=cmd_extract)
    p = sub.add_parser("matrix-init")
    p.add_argument("--force", action="store_true")
    p.set_defaults(fn=cmd_matrix_init)
    sub.add_parser("matrix").set_defaults(fn=cmd_matrix)
    p = sub.add_parser("qa")
    p.add_argument("drafts", nargs="*")
    p.add_argument("--strict", action="store_true", help="warnings fail too")
    p.set_defaults(fn=cmd_qa)
    p = sub.add_parser("export")
    p.add_argument("drafts", nargs="*")
    p.add_argument("--name", default="study")
    p.add_argument("--formats", nargs="+", default=["md", "docx", "pdf"], choices=["md", "docx", "pdf", "html"])
    p.add_argument("--skip-qa", action="store_true")
    p.set_defaults(fn=cmd_export)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
